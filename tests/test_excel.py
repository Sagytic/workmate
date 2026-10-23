import csv
import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

from backend.excel_agent import inspect_workbook, inspect_workbooks, repair_workbook

FIXTURES = Path(__file__).resolve().parents[1] / "data" / "fixtures"


class ExcelAgentTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.base = FIXTURES / "baseline.xlsx"

    def tearDown(self):
        self.temp.cleanup()

    def copy_edit(self, source, name, edits):
        book = load_workbook(source)
        for coordinate, value in edits.items():
            book["업무내역"][coordinate] = value
        target = self.root / name
        book.save(target)
        book.close()
        return target

    def inspect(self, current, baseline=None):
        return inspect_workbook(baseline or self.base, current, self.root / "report")

    def test_known_corruptions_match_independent_ground_truth(self):
        truth = json.loads((FIXTURES / "ground_truth.json").read_text(encoding="utf-8"))
        result = self.inspect(FIXTURES / "changed.xlsx")
        actual = sorted((f["cell"], f["rule"]) for f in result["findings"])
        expected = sorted(tuple(x) for x in truth["expected"]["changed"]["findings"])
        self.assertEqual(actual, expected)
        self.assertEqual({f["cell"] for f in result["findings"] if f["repairable"]}, {"F2", "F3", "F4"})
        self.assertFalse(result["recalculated"])

    def test_clean_and_reordered_have_no_false_positives(self):
        for name in ("clean", "reordered"):
            with self.subTest(name=name):
                result = self.inspect(FIXTURES / f"{name}.xlsx")
                self.assertEqual(result["findings"], [])
                self.assertEqual(result["changes"], [])
                self.assertTrue(result["supported_schema"])
                self.assertTrue(result["formula_scope_supported"])
                self.assertEqual(result["formula_scope"]["supported"], 8)

    def test_repair_changes_only_approved_formula_cells_and_preserves_original(self):
        current = FIXTURES / "changed.xlsx"
        original_hash = hashlib.sha256(current.read_bytes()).hexdigest()
        repaired = repair_workbook(self.base, current, self.root / "repair")
        self.assertEqual(repaired["applied_count"], 3)
        self.assertEqual(hashlib.sha256(current.read_bytes()).hexdigest(), original_hash)
        before, after = load_workbook(current), load_workbook(self.root / "repair" / repaired["filename"])
        try:
            self.assertEqual(before.sheetnames, after.sheetnames)
            for ws in before:
                target = after[ws.title]
                self.assertEqual(str(ws.data_validations), str(target.data_validations))
                for row in ws:
                    for cell in row:
                        new = target[cell.coordinate]
                        if cell.coordinate in {"F2", "F3", "F4"}:
                            self.assertEqual(new.value, f"=D{cell.row}*E{cell.row}")
                        else:
                            self.assertEqual(new.value, cell.value)
                        self.assertEqual(new._style, cell._style)
                self.assertEqual(ws.sheet_view.showGridLines, target.sheet_view.showGridLines)
            self.assertTrue(after.calculation.fullCalcOnLoad)
        finally:
            before.close()
            after.close()
        remaining = repaired["remaining"]["findings"]
        self.assertFalse(any(f["repairable"] for f in remaining))
        self.assertTrue(any(f["rule"] == "duplicate_id" for f in remaining))

    def test_reordered_formula_restore_translates_to_actual_row(self):
        current = self.copy_edit(FIXTURES / "reordered.xlsx", "reordered-broken.xlsx", {"F5": 0})
        result = self.inspect(current)
        self.assertEqual(len(result["findings"]), 1)
        self.assertEqual(result["findings"][0]["before"], "=D2*E2")
        self.assertEqual(result["findings"][0]["suggested"], "=D5*E5")
        repaired = repair_workbook(self.base, current, self.root / "repair")
        self.assertEqual(repaired["applied"][0]["cell"], "F5")

    def test_duplicate_current_key_never_repairs_either_row(self):
        current = self.copy_edit(self.base, "duplicates.xlsx", {"A3": "WK-001", "F2": 1, "F3": 2})
        result = self.inspect(current)
        self.assertEqual(result["repairable_count"], 0)
        self.assertEqual(sum(f["rule"] == "duplicate_id" for f in result["findings"]), 2)
        repaired = repair_workbook(self.base, current, self.root / "repair")
        self.assertEqual(repaired["applied"], [])

    def test_duplicate_baseline_key_never_repairs(self):
        baseline = self.copy_edit(self.base, "base-duplicates.xlsx", {"A3": "WK-001"})
        current = self.copy_edit(self.base, "current.xlsx", {"F2": 1})
        result = self.inspect(current, baseline)
        self.assertEqual(result["repairable_count"], 0)
        self.assertTrue(any(f["rule"] == "baseline_duplicate_id" for f in result["findings"]))

    def test_unsupported_headers_are_explicitly_reported(self):
        current = self.copy_edit(self.base, "unknown.xlsx", {"A1": "record"})
        result = self.inspect(current)
        self.assertFalse(result["supported_schema"])
        self.assertTrue(any(f["rule"] == "unsupported_schema" for f in result["findings"]))
        self.assertEqual(result["repairable_count"], 0)

    def test_cross_row_formula_not_auto_repaired(self):
        baseline = self.copy_edit(self.base, "cross-row-base.xlsx", {"F2": "=D3*E2"})
        current = self.copy_edit(self.base, "cross-row-current.xlsx", {"F2": 0})
        result = self.inspect(current, baseline)
        finding = next(f for f in result["findings"] if f["cell"] == "F2")
        self.assertFalse(finding["repairable"])
        self.assertIsNone(finding["suggested"])

    def test_zero_input_valid_but_missing_input_reported(self):
        current = self.copy_edit(self.base, "zero.xlsx", {"D2": 0, "E3": None})
        result = self.inspect(current)
        self.assertFalse(any(f["cell"] == "D2" for f in result["findings"]))
        self.assertTrue(any(f["cell"] == "E3" and f["rule"] == "required_value" for f in result["findings"]))

    def test_formula_error_text_inside_title_not_reported_as_excel_error(self):
        current = self.copy_edit(self.base, "text.xlsx", {"B2": "#REF! 오류 수정 업무"})
        self.assertEqual(self.inspect(current)["findings"], [])

    def test_quoted_error_string_in_formula_is_not_broken_reference(self):
        baseline = self.copy_edit(self.base, "quoted-base.xlsx", {"F2": '=IF(D2=0,"#REF!",D2*E2)'})
        self.assertEqual(self.inspect(baseline, baseline)["findings"], [])

    def test_malformed_baseline_formula_is_not_repairable(self):
        baseline = self.copy_edit(self.base, "malformed-base.xlsx", {"F2": "=SUM(D2:E2"})
        current = self.copy_edit(self.base, "malformed-current.xlsx", {"F2": 0})
        self.assertEqual(self.inspect(current, baseline)["repairable_count"], 0)

    def test_invalid_expression_or_self_containing_range_is_not_supported_or_restored(self):
        formulas = ("=D2+", "=D2*/E2", "=D2 E2", "=IF(D2)", "=ROUND(D2)", "=ABS(D2,E2)", "=SUM(D2:H2)", "=SUM(D2:E2:G2)", "=SUM({1,2})+D2")
        for index, formula in enumerate(formulas):
            with self.subTest(formula=formula):
                baseline = self.copy_edit(self.base, f"invalid-base-{index}.xlsx", {"F2": formula})
                unchanged = self.inspect(baseline, baseline)
                self.assertFalse(unchanged["formula_scope_supported"])
                current = self.copy_edit(baseline, f"invalid-current-{index}.xlsx", {"F2": 0})
                result = self.inspect(current, baseline)
                self.assertEqual(result["repairable_count"], 0)

    def test_unchanged_unsupported_formulas_do_not_report_false_clean(self):
        for index, formula in enumerate(("=1/0", "=D3*E2", "='[other.xlsx]Sheet1'!D2", "=INDIRECT(\"D2\")", "=SUM(D2:E2")):
            with self.subTest(formula=formula):
                current = self.copy_edit(self.base, f"unsupported-{index}.xlsx", {"F2": formula})
                result = self.inspect(current, current)
                self.assertTrue(result["supported_schema"])
                self.assertFalse(result["formula_scope_supported"])
                self.assertEqual(result["formula_scope"]["unverified"], 1)
                self.assertEqual(result["changes"], [])
                self.assertTrue(any(f["rule"] == "formula_scope_unverified" and f["cell"] == "F2" for f in result["findings"]))

    def test_array_and_data_table_formulas_are_explicitly_unverified(self):
        for index, value in enumerate((ArrayFormula(ref="F2", text="=SUM(D2:E2)"), DataTableFormula(ref="F2", r1="D2"))):
            with self.subTest(formula_type=type(value).__name__):
                current = self.copy_edit(self.base, f"formula-type-{index}.xlsx", {"F2": value})
                result = self.inspect(current, current)
                self.assertFalse(result["formula_scope_supported"])
                self.assertTrue(any(f["rule"] == "formula_scope_unverified" for f in result["findings"]))
                self.assertEqual(result["repairable_count"], 0)
                self.assertEqual(result["changes"], [])

    def test_same_row_sum_and_arithmetic_remain_in_supported_static_scope(self):
        for index, formula in enumerate(("=SUM(D2:E2)", "=D2*E2", "=ROUND(D2*E2,0)", '=IF(D2=0,"#REF!",D2*E2)', "=SUM(D2, E2)", "=-(D2+E2)*10%")):
            with self.subTest(formula=formula):
                current = self.copy_edit(self.base, f"supported-{index}.xlsx", {"F2": formula})
                result = self.inspect(current, current)
                self.assertTrue(result["formula_scope_supported"])
                self.assertEqual(result["findings"], [])

    def test_unchanged_broken_reference_is_still_an_error(self):
        current = self.copy_edit(self.base, "same-error.xlsx", {"F2": "=D2*#REF!"})
        result = self.inspect(current, current)
        self.assertEqual(result["changes"], [])
        self.assertTrue(any(f["rule"] == "broken_reference" for f in result["findings"]))

    def test_html_report_escapes_untrusted_workbook_values(self):
        current = self.copy_edit(self.base, "html.xlsx", {"G2": "<script>alert('bad')</script>"})
        self.inspect(current)
        html = (self.root / "report" / "inspection.html").read_text(encoding="utf-8")
        self.assertNotIn("<script>", html)
        self.assertIn("&lt;script&gt;", html)

    def test_recognized_id_aliases_enable_unique_row_comparison(self):
        for index, header in enumerate(("id", "요청ID", "항목ID", "코드", "request_id")):
            with self.subTest(header=header):
                baseline = self.copy_edit(self.base, f"alias-base-{index}.xlsx", {"A1": header})
                current = self.copy_edit(FIXTURES / "reordered.xlsx", f"alias-current-{index}.xlsx", {"A1": header})
                result = self.inspect(current, baseline)
                self.assertEqual(result["findings"], [])
                self.assertEqual(result["changes"], [])
                self.assertEqual(result["coverage"][0]["matched_unique_rows"], 8)
                self.assertEqual(result["coverage"][0]["key_header"], header)
                self.assertTrue(result["supported_schema"])

    def test_multiple_key_alias_columns_are_ambiguous_and_never_repaired(self):
        current = self.copy_edit(self.base, "ambiguous-keys.xlsx", {"H1": "코드", "H2": "A-001", "F2": 0})
        result = self.inspect(current)
        self.assertEqual(result["repairable_count"], 0)
        self.assertFalse(result["supported_schema"])
        self.assertEqual(result["verification_status"], "unverified")
        self.assertTrue(any(f["rule"] == "unsupported_schema" for f in result["findings"]))

    def test_partial_business_schema_keeps_key_comparison_but_reports_coverage(self):
        baseline = self.copy_edit(self.base, "partial-base.xlsx", {"G1": "설명"})
        current = self.copy_edit(baseline, "partial-current.xlsx", {"B2": "changed title"})
        result = self.inspect(current, baseline)
        self.assertFalse(result["supported_schema"])
        self.assertEqual(result["support_coverage"]["keyed_comparison_sheets"], 1)
        self.assertEqual(result["support_coverage"]["business_rule_sheets"], 0)
        self.assertEqual(result["coverage"][0]["missing_business_headers"], ["status"])
        self.assertEqual(result["counts"]["changes"], 1)
        self.assertTrue(any(f["rule"] == "business_rules_unavailable" for f in result["findings"]))

    def test_csv_exports_have_bom_and_neutralize_formula_strings(self):
        payloads = {
            "B2": '=HYPERLINK("https://example.invalid","x")', "B3": "+SUM(1,2)",
            "B4": "-1+1", "B5": "@SUM(1,2)", "B6": "\t=1+1", "B7": " \r=2+2",
            "G8": "@SUM(8,9)", "D9": -1,
        }
        current = self.copy_edit(self.base, "=untrusted.xlsx", payloads)
        result = self.inspect(current)
        for name in ("findings.csv", "changes.csv"):
            content = (self.root / "report" / name).read_bytes()
            self.assertTrue(content.startswith(b"\xef\xbb\xbf"))
            self.assertTrue(any(a["name"] == name for a in result["artifacts"]))
        with (self.root / "report" / "changes.csv").open(encoding="utf-8-sig", newline="") as stream:
            rows = {row["cell"]: row for row in csv.DictReader(stream) if row["type"] == "cell_changed"}
        for cell, payload in payloads.items():
            self.assertEqual(rows[cell]["after"], "'" + payload if isinstance(payload, str) else str(payload))
            self.assertEqual(rows[cell]["filename"], "'=untrusted.xlsx")
        with (self.root / "report" / "findings.csv").open(encoding="utf-8-sig", newline="") as stream:
            findings = list(csv.DictReader(stream))
        self.assertTrue(any(row["cell"] == "G8" and row["after"] == "'@SUM(8,9)" for row in findings))
        self.assertEqual(next(c["after"] for c in result["changes"] if c.get("cell") == "B3"), "+SUM(1,2)")

    def test_csv_escapes_formula_before_and_suggested_without_altering_json(self):
        result = self.inspect(FIXTURES / "changed.xlsx")
        with (self.root / "report" / "findings.csv").open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        restored = next(row for row in rows if row["rule"] == "formula_replaced")
        self.assertEqual(restored["before"], "'=D2*E2")
        self.assertEqual(restored["suggested"], "'=D2*E2")
        self.assertEqual(next(f["suggested"] for f in result["findings"] if f["rule"] == "formula_replaced"), "=D2*E2")

    def test_batch_aggregates_results_and_returns_resolvable_prefixed_artifacts(self):
        output = self.root / "batch"
        result = inspect_workbooks(self.base, [FIXTURES / "clean.xlsx", FIXTURES / "changed.xlsx", FIXTURES / "reordered.xlsx"], output)
        self.assertEqual(result["status"], "completed")
        self.assertEqual(result["batch_counts"], {"total_files": 3, "completed_files": 3, "failed_files": 0, "unverified_files": 0})
        self.assertEqual(result["counts"], result["batches"][1]["counts"])
        self.assertEqual(result["suggested_repairs"], 3)
        self.assertEqual(result["repairable_count"], 0)
        self.assertFalse(result["batch_repair_supported"])
        self.assertEqual(len({f["id"] for f in result["findings"]}), len(result["findings"]))
        self.assertTrue(all(f["filename"] == "changed.xlsx" for f in result["findings"]))
        for artifact in result["artifacts"]:
            self.assertTrue((output / artifact["name"]).is_file())
        self.assertTrue(all(a["name"].startswith("case-02/") for a in result["batches"][1]["artifacts"]))
        saved = json.loads((output / "inspection.json").read_text(encoding="utf-8"))
        self.assertEqual(saved["artifacts"], result["artifacts"])
        html = (output / "inspection.html").read_text(encoding="utf-8")
        self.assertIn("case-02/inspection.html", html)

    def test_batch_isolates_corrupt_and_missing_files_without_false_pass(self):
        corrupt = self.root / "corrupt.xlsx"
        corrupt.write_bytes(b"not an XLSX zip archive")
        output = self.root / "batch"
        result = inspect_workbooks(self.base, [corrupt, FIXTURES / "clean.xlsx", self.root / "missing.xlsx", FIXTURES / "changed.xlsx"], output)
        self.assertEqual(result["status"], "partial")
        self.assertEqual(result["verification_status"], "unverified")
        self.assertEqual(result["batch_counts"]["failed_files"], 2)
        self.assertFalse(result["supported_schema"])
        self.assertFalse(result["formula_scope_supported"])
        self.assertEqual(result["batches"][1]["findings"], [])
        self.assertEqual(result["batches"][3]["repairable_count"], 3)
        failures = [f for f in result["findings"] if f["rule"] == "inspection_failed"]
        self.assertEqual(len(failures), 2)
        self.assertTrue(all(f["severity"] == "error" and not f["repairable"] for f in failures))
        for case_id in ("case-01", "case-03"):
            failed = json.loads((output / case_id / "inspection.json").read_text(encoding="utf-8"))
            self.assertEqual(failed["status"], "failed")
            self.assertFalse(failed["support_coverage"]["file_inspected"])

    def test_batch_unreadable_baseline_marks_every_file_unverified(self):
        result = inspect_workbooks(self.root / "missing-base.xlsx", [FIXTURES / "clean.xlsx", FIXTURES / "changed.xlsx"], self.root / "batch")
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["batch_counts"]["failed_files"], 2)
        self.assertEqual(result["counts"]["errors"], 2)
        self.assertEqual(result["formula_scope"]["total"], 0)
        self.assertFalse(result["formula_scope_supported"])
        self.assertIsNone(result["baseline_sha256"])

    def test_batch_rejects_empty_and_oversized_input_before_writing(self):
        output = self.root / "batch"
        for currents in ([], [self.base] * 11, str(self.base)):
            with self.subTest(currents=str(currents)[:100]):
                with self.assertRaises(ValueError):
                    inspect_workbooks(self.base, currents, output)
        self.assertFalse(output.exists())

    def test_batch_same_filename_inputs_have_separate_artifact_directories(self):
        result = inspect_workbooks(self.base, [FIXTURES / "changed.xlsx", FIXTURES / "changed.xlsx"], self.root / "batch")
        names = [a["name"] for a in result["artifacts"]]
        self.assertEqual(len(names), len(set(names)))
        ids = [f["id"] for f in result["findings"]]
        self.assertEqual(len(ids), len(set(ids)))

    def test_repair_cannot_overwrite_input(self):
        current = self.copy_edit(self.base, "repaired.xlsx", {"F2": 0})
        original = current.read_bytes()
        with self.assertRaisesRegex(ValueError, "원본"):
            repair_workbook(self.base, current, self.root)
        self.assertEqual(current.read_bytes(), original)


if __name__ == "__main__":
    unittest.main()

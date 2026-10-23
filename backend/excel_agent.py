"""Deterministic XLSX comparison and cautious, baseline-based formula restoration.

No LLM claims, no Excel calculation engine, and no original workbook writes.
The operational demo schema is explicit; unsupported sheets are never called clean.
"""
from __future__ import annotations

from collections import defaultdict
from datetime import date, datetime
from hashlib import sha256
from html import escape
import csv
import json
import re
import shutil
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.formula.tokenizer import TokenizerError
from openpyxl.formula.translate import Translator
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from openpyxl.workbook.properties import CalcProperties
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

RULE_VERSION = "workmate-excel-1.0"
MAX_CELLS = 200_000
MAX_UNCOMPRESSED = 100 * 1024 * 1024
MAX_BATCH_FILES = 10
ALIASES = {
    "id": {"id", "고유id", "고유번호", "관리번호", "요청id", "항목id", "코드", "request_id", "item_id"},
    "title": {"title", "제목", "항목", "업무명"},
    "owner": {"owner", "담당자"},
    "quantity": {"quantity", "수량"},
    "unit_price": {"unit_price", "단가"},
    "total": {"total", "합계", "금액"},
    "status": {"status", "상태"},
}
ALLOWED_STATUSES = {"대기", "진행", "완료", "보류"}
SAFE_FUNCTIONS = {"SUM", "PRODUCT", "ROUND", "ROUNDUP", "ROUNDDOWN", "ABS", "MIN", "MAX", "AVERAGE", "IF"}
ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!", "#SPILL!", "#CALC!"}


def _json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    if isinstance(value, (ArrayFormula, DataTableFormula)):
        return json.dumps({"type": type(value).__name__, **dict(value), "text": getattr(value, "text", None)}, ensure_ascii=False, sort_keys=True)
    return str(value)


def _hash(path: Path) -> str:
    return sha256(path.read_bytes()).hexdigest()


def _load(path: Path):
    path = Path(path)
    if path.suffix.lower() != ".xlsx":
        raise ValueError(".xlsx 파일만 지원합니다. 매크로 파일은 수정하지 않습니다.")
    with ZipFile(path) as archive:
        if sum(i.file_size for i in archive.infolist()) > MAX_UNCOMPRESSED:
            raise ValueError("압축 해제 크기가 100MB를 넘는 파일은 지원하지 않습니다.")
    book = load_workbook(path, data_only=False, keep_links=True)
    if sum(ws.max_row * ws.max_column for ws in book) > MAX_CELLS:
        book.close()
        raise ValueError("검사 범위가 20만 셀을 넘습니다. 더 작은 파일로 나누어 주세요.")
    return book


def _header(ws):
    for row in range(1, min(ws.max_row, 20) + 1):
        names = [str(ws.cell(row, c).value or "").strip().lower() for c in range(1, ws.max_column + 1)]
        mapped = {}
        populated_names = [name for name in names if name]
        ambiguous = len(populated_names) != len(set(populated_names))
        for canonical, aliases in ALIASES.items():
            indexes = [i + 1 for i, name in enumerate(names) if name in aliases]
            if len(indexes) > 1:
                ambiguous = True
            elif indexes:
                mapped[canonical] = indexes[0]
        if "id" in mapped:
            return row, mapped, ambiguous
    return None, {}, False


def _key(value):
    # Do not conflate numeric IDs with text IDs or silently trim identifiers.
    if value is None or (isinstance(value, str) and not value.strip()):
        return None
    return type(value).__name__, str(value)


def _rows(ws, header_row, id_column):
    rows, keys = [], defaultdict(list)
    for row in range(header_row + 1, ws.max_row + 1):
        if not any(ws.cell(row, c).value is not None for c in range(1, ws.max_column + 1)):
            continue
        rows.append(row)
        key = _key(ws.cell(row, id_column).value)
        if key is not None and ws.cell(row, id_column).data_type != "f":
            keys[key].append(row)
    return rows, keys


def _formula(cell):
    return cell.data_type == "f" and isinstance(cell.value, str)


def _broken_reference(cell):
    if not _formula(cell):
        return cell.data_type == "e" and cell.value == "#REF!"
    try:
        return any(t.type == "OPERAND" and t.subtype == "ERROR" and t.value.upper() == "#REF!" for t in Tokenizer(cell.value).items)
    except (TokenizerError, ValueError):
        return False  # A malformed expression is not proof of a broken reference.


def _row_formula(formula: str, source: str, target: str) -> str | None:
    """Restore only bounded same-row formulas; cross-row reorder semantics are ambiguous."""
    try:
        source_col, source_row = coordinate_from_string(source)
        tokens = Tokenizer(formula).items
        references = 0
        expect_operand = True
        stack = []
        for token in tokens:
            if token.type == "WHITE-SPACE":
                continue
            if token.type in {"FUNC", "PAREN"} and token.subtype == "OPEN":
                if not expect_operand:
                    return None
                function = token.value[:-1].upper() if token.type == "FUNC" else None
                if function is not None and function not in SAFE_FUNCTIONS:
                    return None
                stack.append({"type": token.type, "function": function, "arguments": 1})
            elif token.type in {"FUNC", "PAREN"} and token.subtype == "CLOSE":
                if expect_operand or not stack or stack[-1]["type"] != token.type:
                    return None
                context = stack.pop()
                bounds = {"ROUND": (2, 2), "ROUNDUP": (2, 2), "ROUNDDOWN": (2, 2), "ABS": (1, 1), "IF": (2, 3)}
                minimum, maximum = bounds.get(context["function"], (1, 255))
                if not minimum <= context["arguments"] <= maximum:
                    return None
                expect_operand = False
            elif token.type == "OPERAND":
                if not expect_operand or token.subtype not in {"RANGE", "NUMBER", "TEXT", "LOGICAL"}:
                    return None
                if token.subtype == "RANGE":
                    addresses = token.value.split(":")
                    if len(addresses) > 2:
                        return None
                    columns = []
                    for address in addresses:
                        match = re.fullmatch(r"\$?([A-Za-z]{1,3})([1-9][0-9]*)", address)
                        if not match or int(match[2]) != source_row:
                            return None
                        column = column_index_from_string(match[1])
                        if column > 16384:
                            return None
                        columns.append(column)
                    if min(columns) <= column_index_from_string(source_col) <= max(columns):
                        return None  # A range may contain the formula cell between its endpoints.
                    references += len(addresses)
                expect_operand = False
            elif token.type == "OPERATOR-PREFIX" and token.value in {"+", "-"}:
                if not expect_operand:
                    return None
            elif token.type == "OPERATOR-POSTFIX" and token.value == "%":
                if expect_operand:
                    return None
            elif token.type == "OPERATOR-INFIX" and token.value in {"+", "-", "*", "/", "^", "&", "=", "<>", "<=", ">=", "<", ">"}:
                if expect_operand:
                    return None
                expect_operand = True
            elif token.type == "SEP" and token.subtype == "ARG":
                if expect_operand or not stack or stack[-1]["type"] != "FUNC":
                    return None
                stack[-1]["arguments"] += 1
                expect_operand = True
            else:
                # Array literals, reference unions, and other grammar are unverified.
                return None
        if not references or stack or expect_operand:
            return None
        return Translator(formula, origin=source).translate_formula(target)
    except (ValueError, IndexError, TypeError, TokenizerError):
        return None


def _validation_signature(ws, coordinate):
    rules = []
    for rule in ws.data_validations.dataValidation:
        if coordinate in rule:
            rules.append((rule.type, rule.operator, rule.formula1, rule.formula2, rule.allowBlank))
    return sorted(rules, key=repr)


def _csv_value(value: Any) -> Any:
    """Protect exported text from being executed when Excel opens the CSV."""
    value = _json_value(value) if not isinstance(value, (dict, list, tuple)) else json.dumps(value, ensure_ascii=False)
    if isinstance(value, str):
        significant = value.lstrip(" \t\r\n")
        if (significant and significant[0] in "=+@-") or value.startswith(("\t", "\r", "\n")):
            return "'" + value
    return value


def _write_csv(path: Path, records: list[dict], columns: list[str]):
    with path.open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(stream, fieldnames=columns, extrasaction="ignore")
        writer.writeheader()
        for record in records:
            writer.writerow({column: _csv_value(record.get(column)) for column in columns})


def _write_report(result: dict, output_dir: Path, extra_artifacts: list[dict] | None = None):
    output_dir.mkdir(parents=True, exist_ok=True)
    result["artifacts"] = [
        {"name": "inspection.json", "kind": "json"}, {"name": "inspection.html", "kind": "html"},
        {"name": "findings.csv", "kind": "csv"}, {"name": "changes.csv", "kind": "csv"},
        *(extra_artifacts or []),
    ]
    export_findings = [{"filename": result.get("current_filename", ""), **record} for record in result["findings"]]
    export_changes = [{"filename": result.get("current_filename", ""), **record} for record in result["changes"]]
    _write_csv(output_dir / "findings.csv", export_findings, ["filename", "id", "severity", "sheet", "cell", "rule", "message", "before", "after", "suggested", "repairable"])
    _write_csv(output_dir / "changes.csv", export_changes, ["filename", "type", "sheet", "cell", "baseline_cell", "key", "field", "before", "after", "expected_at_current_row", "row", "baseline_row"])
    (output_dir / "inspection.json").write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")

    def h(value):
        return escape(str(value if value is not None else "—"))

    def location(record):
        prefix = f"{h(record['filename'])}<br>" if record.get("filename") else ""
        return prefix + h(record.get("sheet", "")) + ("!" + h(record["cell"]) if record.get("cell") else "")

    finding_rows = "".join(f"<tr><td><span class='badge'>{h('오류' if f['severity'] == 'error' else '검토')}</span><small>{h(f['rule'])}</small></td><td>{location(f)}</td><td>{h(f['message'])}</td><td><code>{h(f.get('before'))}</code></td><td><code>{h(f.get('after'))}</code></td><td><code>{h(f.get('suggested'))}</code></td></tr>" for f in result["findings"])
    finding_rows = finding_rows or "<tr><td colspan='6' class='empty'>지원하는 정적 검사 범위에서 발견 사항이 없습니다.</td></tr>"
    change_rows = "".join(f"<tr><td>{h(c['type'])}</td><td>{location(c)}</td><td>{h(c.get('key'))}<small>{h(c.get('field', ''))}</small></td><td><code>{h(c.get('before'))}</code></td><td><code>{h(c.get('after'))}</code></td></tr>" for c in result["changes"])
    change_rows = change_rows or "<tr><td colspan='5' class='empty'>수행한 비교 범위에서 변경 사항이 없습니다.</td></tr>"
    coverage_labels = {"business_rules": "ID 비교 + 업무 규칙", "keyed_comparison": "ID 비교 · 업무 규칙 미지원", "static_only": "정적 오류 검색만 수행", "failed": "검사 실패 · 미검증"}
    coverage_rows = "".join(f"<tr><td>{location(c)}</td><td>{h(coverage_labels.get(c['status'], c['status']))}</td><td>{h(c.get('rows'))}</td><td>{h(c.get('matched_unique_rows'))}</td><td>{h(', '.join(c.get('missing_business_headers', [])) or '—')}</td></tr>" for c in result.get("coverage", []))
    coverage_rows = coverage_rows or "<tr><td colspan='5' class='empty'>검증된 시트 범위가 없습니다.</td></tr>"
    limitations = "".join(f"<li>{h(item)}</li>" for item in result["limitations"])
    batch_section = ""
    if result.get("batches"):
        batch_rows = "".join(f"<tr><td><a href='{h(b['case_id'])}/inspection.html'>{h(b['filename'])}</a></td><td>{h('실패 · 미검증' if b['status'] == 'failed' else '검사 완료')}</td><td>{h(b['summary'])}</td></tr>" for b in result["batches"])
        batch_section = f"<section><h2>파일별 결과</h2><div class='scroll'><table><thead><tr><th>파일</th><th>실행 상태</th><th>결과</th></tr></thead><tbody>{batch_rows}</tbody></table></div></section>"
    scope = result.get("formula_scope", {})
    scope_text = f"수식 {scope.get('total', 0)}개 중 정적 지원 {scope.get('supported', 0)}개 · 알려진 오류 {scope.get('known_errors', 0)}개 · 미검증 {scope.get('unverified', 0)}개"
    filename_text = f"기준: {h(result.get('baseline_filename'))}" + (f" · 비교: {h(result['current_filename'])}" if result.get("current_filename") else "")
    counts = result["counts"]
    (output_dir / "inspection.html").write_text(f"""<!doctype html>
<html lang="ko"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Excel 검사 결과</title>
<style>
*{{box-sizing:border-box}}body{{font:14px/1.6 -apple-system,BlinkMacSystemFont,'Segoe UI',sans-serif;background:#f5f5f4;color:#242627;margin:0;padding:40px 24px}}main{{max-width:1440px;margin:auto}}header,section{{background:#fff;border:1px solid #e4e5e2;border-radius:14px;margin-bottom:18px;padding:26px}}h1{{font-size:28px;letter-spacing:-1px;margin:4px 0 8px}}h2{{font-size:18px;margin:0 0 16px}}p{{margin:10px 0}}.eyebrow{{font-size:11px;letter-spacing:2px;color:#6f746f;font-weight:700}}.muted,small{{color:#727772}}small{{display:block;font-size:12px;overflow-wrap:anywhere}}.cards{{display:grid;grid-template-columns:repeat(3,1fr);gap:12px;margin:20px 0}}.card{{border:1px solid #e4e5e2;border-radius:10px;padding:16px}}.card strong{{display:block;font-size:29px;line-height:1.3}}.note{{border-left:3px solid #a6ad9e;background:#f5f6f3;padding:13px 16px;font-size:13px}}.downloads{{display:flex;gap:10px;flex-wrap:wrap;margin-top:16px}}.downloads a{{border:1px solid #d9ddda;border-radius:7px;padding:7px 12px;text-decoration:none;color:#37483c}}a{{color:#395a45}}.scroll{{overflow:auto}}table{{border-collapse:collapse;width:100%;font-size:13px}}th,td{{padding:13px 12px;text-align:left;border-bottom:1px solid #eaede8;vertical-align:top;min-width:90px}}th{{font-weight:600;color:#5d655e;background:#f8f9f7;white-space:nowrap}}td{{max-width:380px;overflow-wrap:anywhere}}code{{font:12px/1.6 ui-monospace,Consolas,monospace;white-space:pre-wrap;overflow-wrap:anywhere}}.badge{{display:inline-block;background:#eef0eb;border-radius:5px;padding:2px 7px;margin-bottom:5px}}.empty{{padding:28px;color:#717870}}footer{{padding:10px 4px 30px}}li{{margin:7px 0}}@media(max-width:640px){{body{{padding:16px 10px}}header,section{{padding:18px}}h1{{font-size:24px}}.cards{{gap:6px}}.card{{padding:12px}}}}
</style></head><body><main><header><div class="eyebrow">WORKMATE / EXCEL REVIEW</div><h1>Excel 검사 결과</h1><p class="muted">{filename_text}</p><p>{h(result['summary'])}</p><div class="cards"><div class="card"><span class="muted">오류</span><strong>{counts['errors']}</strong></div><div class="card"><span class="muted">검토 필요</span><strong>{counts['warnings']}</strong></div><div class="card"><span class="muted">변경 사항</span><strong>{counts['changes']}</strong></div></div><p class="note">Excel 재계산을 실행하지 않았습니다. 기준 파일과 지원하는 정적 규칙에 따른 검사이며 업무상 정답을 보증하지 않습니다. 미지원 범위와 검사 실패 파일은 검증 완료로 처리하지 않습니다.</p><nav class="downloads"><a href="findings.csv">발견 사항 CSV</a><a href="changes.csv">변경 목록 CSV</a><a href="inspection.json">전체 결과 JSON</a></nav><small>CSV는 UTF-8 BOM 형식이며 수식으로 해석될 수 있는 문자열 앞에 작은따옴표를 붙여 텍스트로 내보냅니다.</small></header>
{batch_section}<section><h2>발견 사항</h2><div class="scroll"><table><thead><tr><th>구분 / 규칙</th><th>위치</th><th>설명</th><th>기준</th><th>현재</th><th>복원 제안</th></tr></thead><tbody>{finding_rows}</tbody></table></div></section>
<section><h2>변경 목록</h2><div class="scroll"><table><thead><tr><th>변경 유형</th><th>위치</th><th>ID / 항목</th><th>기준</th><th>현재</th></tr></thead><tbody>{change_rows}</tbody></table></div></section>
<section><h2>지원 범위와 검사 한계</h2><p>{h(scope_text)}</p><p class="muted">업무 규칙에 필요한 헤더: ID, title, owner, quantity, unit_price, total, status 또는 지원하는 별칭. 행 비교는 양쪽에서 유일한 ID에만 적용합니다.</p><div class="scroll"><table><thead><tr><th>시트</th><th>적용 범위</th><th>현재 행</th><th>고유 ID 대응 행</th><th>부족한 업무 헤더</th></tr></thead><tbody>{coverage_rows}</tbody></table></div><ul>{limitations}</ul></section>
<footer><small>규칙 버전: {h(RULE_VERSION)}<br>기준 SHA-256: {h(result.get('baseline_sha256'))}<br>현재 SHA-256: {h(result.get('current_sha256', '파일별 결과 참조'))}</small></footer></main></body></html>""", encoding="utf-8")


def inspect_workbook(baseline: Path, current: Path, output_dir: Path) -> dict:
    baseline, current, output_dir = map(Path, (baseline, current, output_dir))
    base, now = _load(baseline), None
    findings, changes, coverage = [], [], []
    formula_scope = {"method": "same_row_reference_static", "total": 0, "supported": 0, "known_errors": 0, "unverified": 0}
    limitations = [
        "Excel 계산 엔진을 실행하지 않았습니다. 저장된 계산값의 최신 여부, 모든 함수의 계산 결과, 순환 참조 전체를 검증하지 않습니다.",
        "행 비교는 동일한 시트명과 양쪽에서 유일한 ID를 사용합니다. ID가 없거나 중복이면 해당 행의 자동 복원을 차단합니다.",
        "필수값·허용 상태·음수 검사는 ID/title/owner/quantity/unit_price/total/status 업무 스키마가 모두 있을 때만 적용합니다.",
        "복원은 기준 파일의 같은 행 내부 수식으로 제한합니다. 기준 파일의 업무상 정확성은 별도 확인이 필요합니다.",
        "수식 정적 검사 범위는 같은 행의 상대 행 참조와 제한된 함수(SUM, PRODUCT, ROUND 계열, ABS, MIN, MAX, AVERAGE, IF)입니다. 상수만 있는 식·교차행·외부/이름 참조·배열/데이터 표 수식은 미검증 경고로 남깁니다.",
    ]
    def finding(rule, severity, sheet, cell, message, before=None, after=None, suggested=None, repairable=False):
        findings.append({"id": f"F{len(findings) + 1:03}", "rule": rule, "severity": severity, "sheet": sheet, "cell": cell, "message": message, "before": _json_value(before), "after": _json_value(after), "suggested": suggested, "repairable": bool(repairable)})
    try:
        now = _load(current)
        for ws in now:
            for row in ws:
                for cell in row:
                    broken = _broken_reference(cell)
                    if broken or (cell.data_type == "e" and cell.value in ERRORS):
                        finding("broken_reference" if broken else "stored_error", "error", ws.title, cell.coordinate, "깨진 참조 또는 저장된 오류가 있습니다. 재계산 결과는 확인하지 않았습니다.", after=cell.value)
                    if cell.data_type == "f":
                        formula_scope["total"] += 1
                        if broken:
                            formula_scope["known_errors"] += 1
                        elif _formula(cell) and _row_formula(cell.value, cell.coordinate, cell.coordinate) is not None:
                            formula_scope["supported"] += 1
                        else:
                            formula_scope["unverified"] += 1
                            finding("formula_scope_unverified", "warning", ws.title, cell.coordinate, "지원하는 수식 정적 검사 범위를 벗어났습니다. 기준과 같아도 유효성을 확인한 것으로 처리하지 않습니다. Excel에서 계산·검토하세요.", after=cell.value)
            if ws.title not in base.sheetnames:
                finding("sheet_added", "warning", ws.title, "", "기준 파일에 없는 시트입니다. 정적 오류 검색만 수행했습니다.")
                changes.append({"type": "sheet_added", "sheet": ws.title})
                coverage.append({"sheet": ws.title, "status": "static_only"})
                continue
            bw = base[ws.title]
            bh, bc, ba = _header(bw)
            ch, cc, ca = _header(ws)
            if bh is None or ch is None or ba or ca:
                finding("unsupported_schema", "warning", ws.title, "", "유일한 ID 헤더를 찾을 수 없거나 헤더가 중복입니다. 행 비교·업무 규칙·자동 복원을 수행하지 않았습니다.")
                coverage.append({"sheet": ws.title, "status": "static_only"})
                continue
            brs, bkeys = _rows(bw, bh, bc["id"])
            crs, ckeys = _rows(ws, ch, cc["id"])
            business = set(ALIASES).issubset(cc)
            layout_same = bw.max_column == ws.max_column and all(bw.cell(bh, c).value == ws.cell(ch, c).value for c in range(1, ws.max_column + 1))
            coverage.append({
                "sheet": ws.title, "status": "business_rules" if business else "keyed_comparison",
                "rows": len(crs), "baseline_rows": len(brs), "layout_same": layout_same,
                "matched_unique_rows": sum(len(rows) == 1 and len(bkeys.get(key, [])) == 1 for key, rows in ckeys.items()),
                "ambiguous_or_missing_key_rows": sum(_key(ws.cell(row, cc["id"]).value) is None or ws.cell(row, cc["id"]).data_type == "f" or len(ckeys.get(_key(ws.cell(row, cc["id"]).value), [])) != 1 or len(bkeys.get(_key(ws.cell(row, cc["id"]).value), [])) > 1 for row in crs),
                "missing_business_headers": sorted(set(ALIASES) - set(cc)),
                "key_header": _json_value(ws.cell(ch, cc["id"]).value),
            })
            if not business:
                finding("business_rules_unavailable", "warning", ws.title, "", "지원 업무 헤더 일부가 없어 ID 비교와 정적 수식 검사만 수행했습니다.")
            if not layout_same:
                finding("column_structure_changed", "warning", ws.title, "", "열 구조가 달라 수식 자동 복원을 차단했습니다. 같은 이름의 헤더만 비교합니다.")
            for key, rows in bkeys.items():
                if len(rows) > 1:
                    finding("baseline_duplicate_id", "warning", ws.title, bw.cell(rows[0], bc["id"]).coordinate, f"기준 파일 ID {key[1]}가 중복입니다. 관련 행은 비교·복원하지 않습니다.", before=key[1])
            for row in crs:
                idcell = ws.cell(row, cc["id"])
                key = _key(idcell.value)
                if key is None or idcell.data_type == "f":
                    finding("missing_id" if key is None else "formula_id", "error", ws.title, idcell.coordinate, "비어 있거나 수식인 ID는 행 정렬 키로 사용할 수 없습니다.", after=idcell.value)
                elif len(ckeys.get(key, [])) > 1:
                    finding("duplicate_id", "error", ws.title, idcell.coordinate, f"ID {key[1]}가 중복입니다. 이 ID의 모든 행에서 자동 복원을 차단합니다.", after=idcell.value)
                if business:
                    for field in ("title", "owner", "quantity", "unit_price", "status"):
                        cell = ws.cell(row, cc[field])
                        if cell.value is None or (isinstance(cell.value, str) and not cell.value.strip()):
                            finding("required_value", "error", ws.title, cell.coordinate, f"필수 항목 {field}가 비어 있습니다.", after=cell.value)
                    for field in ("quantity", "unit_price"):
                        cell = ws.cell(row, cc[field])
                        if cell.value is None:
                            continue
                        if isinstance(cell.value, bool) or not isinstance(cell.value, (int, float)):
                            finding("numeric_input", "warning", ws.title, cell.coordinate, f"{field}는 숫자 입력을 기대합니다. 수식 입력의 계산값은 검증하지 않습니다.", after=cell.value)
                        elif cell.value < 0:
                            finding("negative_input", "error", ws.title, cell.coordinate, f"업무 규칙에서 {field}는 0 이상이어야 합니다.", after=cell.value)
                    status = ws.cell(row, cc["status"])
                    if status.value is not None and status.value not in ALLOWED_STATUSES:
                        finding("invalid_status", "error", ws.title, status.coordinate, "허용 상태: 대기, 진행, 완료, 보류", after=status.value)
            for key, rows in ckeys.items():
                if len(rows) != 1 or len(bkeys.get(key, [])) > 1:
                    continue
                if key not in bkeys:
                    changes.append({"type": "row_added", "sheet": ws.title, "key": key[1], "row": rows[0]})
                    continue
                cr, br = rows[0], bkeys[key][0]
                # Compare columns by header identity. Reordering alone is not a changed business value.
                bcols = {str(bw.cell(bh, c).value): c for c in range(1, bw.max_column + 1) if bw.cell(bh, c).value is not None}
                ccols = {str(ws.cell(ch, c).value): c for c in range(1, ws.max_column + 1) if ws.cell(ch, c).value is not None}
                for header in bcols.keys() & ccols.keys():
                    before, after = bw.cell(br, bcols[header]), ws.cell(cr, ccols[header])
                    restored = _row_formula(before.value, before.coordinate, after.coordinate) if _formula(before) and layout_same else None
                    expected = restored if restored is not None else before.value
                    if _json_value(expected) != _json_value(after.value):
                        changes.append({"type": "cell_changed", "sheet": ws.title, "cell": after.coordinate, "baseline_cell": before.coordinate, "key": key[1], "field": header, "before": _json_value(before.value), "after": _json_value(after.value), "expected_at_current_row": _json_value(expected)})
                        if _formula(before):
                            rule = "missing_formula" if after.value is None else "formula_changed" if _formula(after) else "formula_replaced"
                            message = {"missing_formula": "기준에 있던 수식이 비어 있습니다.", "formula_replaced": "기준 수식이 상수 또는 텍스트로 바뀌었습니다.", "formula_changed": "기준과 수식이 다릅니다. 의도한 업무 변경인지 검토하세요."}[rule]
                            finding(rule, "error" if rule != "formula_changed" else "warning", ws.title, after.coordinate, message, before.value, after.value, restored, restored is not None)
                    old_rule = _validation_signature(bw, before.coordinate)
                    new_rule = _validation_signature(ws, after.coordinate)
                    if old_rule != new_rule:
                        changes.append({"type": "validation_changed", "sheet": ws.title, "cell": after.coordinate, "key": key[1], "before": old_rule, "after": new_rule})
                        finding("validation_changed", "warning", ws.title, after.coordinate, "기준과 데이터 유효성 규칙이 다릅니다. 규칙을 자동 복원하지 않습니다.", repr(old_rule), repr(new_rule))
            for key, rows in bkeys.items():
                if len(rows) == 1 and key not in ckeys:
                    changes.append({"type": "row_removed", "sheet": ws.title, "key": key[1], "baseline_row": rows[0]})
        for name in set(base.sheetnames) - set(now.sheetnames):
            finding("sheet_removed", "warning", name, "", "기준 파일의 시트가 삭제되었습니다.")
            changes.append({"type": "sheet_removed", "sheet": name})
        # Stable ordering makes saved reports and tests reproducible.
        findings.sort(key=lambda f: (f["sheet"], f["cell"], f["rule"]))
        for i, f in enumerate(findings, 1):
            f["id"] = f"F{i:03}"
        changes.sort(key=lambda c: (c["sheet"], c.get("cell", ""), c["type"], c.get("key", "")))
        counts = {"errors": sum(f["severity"] == "error" for f in findings), "warnings": sum(f["severity"] == "warning" for f in findings), "changes": len(changes)}
        supported_schema = bool(coverage) and all(c["status"] == "business_rules" for c in coverage)
        result = {"summary": f"오류 {counts['errors']}건 · 검토 {counts['warnings']}건 · 변경 {counts['changes']}건. Excel 재계산 미실행.", "status": "completed", "verification_status": "unverified" if not supported_schema or formula_scope["unverified"] else "findings" if findings else "within_supported_static_scope", "findings": findings, "changes": changes, "counts": counts, "recalculated": False, "rule_version": RULE_VERSION, "coverage": coverage, "supported_schema": supported_schema, "formula_scope": formula_scope, "formula_scope_supported": formula_scope["unverified"] == 0, "limitations": limitations, "baseline_filename": baseline.name, "current_filename": current.name, "baseline_sha256": _hash(baseline), "current_sha256": _hash(current), "repairable_count": sum(f["repairable"] for f in findings)}
        result["support_coverage"] = {
            "expected_business_headers": list(ALIASES), "accepted_key_headers": sorted(ALIASES["id"]),
            "current_sheets": len(now.sheetnames), "baseline_sheets": len(base.sheetnames),
            "keyed_comparison_sheets": sum(c["status"] in {"business_rules", "keyed_comparison"} for c in coverage),
            "business_rule_sheets": sum(c["status"] == "business_rules" for c in coverage),
            "static_only_sheets": sum(c["status"] == "static_only" for c in coverage),
            "matched_unique_rows": sum(c.get("matched_unique_rows", 0) for c in coverage),
            "unverified_formula_cells": formula_scope["unverified"], "calculation_verified": False,
        }
        _write_report(result, output_dir)
        return result
    finally:
        base.close()
        if now is not None:
            now.close()


def inspect_workbooks(baseline: Path, currents: list[Path], output_dir: Path) -> dict:
    """Compare up to ten files independently, preserving failed files as unverified."""
    if isinstance(currents, (str, Path)):
        raise ValueError("비교할 파일 경로를 목록으로 전달하세요.")
    currents = list(currents)
    if not 1 <= len(currents) <= MAX_BATCH_FILES:
        raise ValueError("일괄 검사는 비교 파일 1~10개를 지원합니다.")
    baseline, output_dir = Path(baseline), Path(output_dir)
    batches, findings, changes, coverage, artifacts = [], [], [], [], []
    formula_scope = {"method": "same_row_reference_static", "total": 0, "supported": 0, "known_errors": 0, "unverified": 0}
    for index, current in enumerate(currents, 1):
        current = Path(current)
        case_id = f"case-{index:02}"
        case_dir = output_dir / case_id
        try:
            result = inspect_workbook(baseline, current, case_dir)
        except Exception as exc:
            # One corrupt/unsupported workbook must not hide results from the others.
            message = f"파일을 검사하지 못했습니다. 결과는 미검증입니다. {type(exc).__name__}: {str(exc)[:500]}"
            result = {
                "summary": "검사 실패 · 미검증. 이 파일의 수식과 변경 사항을 확인하지 못했습니다.",
                "status": "failed", "verification_status": "unverified", "error": str(exc)[:500],
                "findings": [{"id": "F001", "rule": "inspection_failed", "severity": "error", "sheet": "", "cell": "", "message": message, "before": None, "after": None, "suggested": None, "repairable": False}],
                "changes": [], "counts": {"errors": 1, "warnings": 0, "changes": 0},
                "coverage": [{"sheet": "", "status": "failed"}], "supported_schema": False,
                "formula_scope": {"method": "same_row_reference_static", "total": 0, "supported": 0, "known_errors": 0, "unverified": 0},
                "formula_scope_supported": False, "support_coverage": {"calculation_verified": False, "file_inspected": False},
                "recalculated": False, "rule_version": RULE_VERSION, "repairable_count": 0,
                "baseline_filename": baseline.name, "current_filename": current.name,
                "baseline_sha256": None, "current_sha256": None,
                "limitations": ["이 파일은 읽기 또는 검사 과정에서 실패했습니다. 0건으로 표시된 변경·수식 수는 확인 결과가 아니라 미집계 상태입니다."],
            }
            _write_report(result, case_dir)
        prefixed_artifacts = [{**artifact, "name": f"{case_id}/{artifact['name']}"} for artifact in result["artifacts"]]
        batches.append({**result, "case_id": case_id, "filename": current.name, "artifacts": prefixed_artifacts})
        artifacts.extend(prefixed_artifacts)
        findings.extend({**f, "id": f"{case_id}-{f['id']}", "filename": current.name, "case_id": case_id} for f in result["findings"])
        changes.extend({**c, "filename": current.name, "case_id": case_id} for c in result["changes"])
        coverage.extend({**c, "filename": current.name, "case_id": case_id} for c in result["coverage"])
        for metric in ("total", "supported", "known_errors", "unverified"):
            formula_scope[metric] += result["formula_scope"][metric]
    failed = sum(b["status"] == "failed" for b in batches)
    counts = {metric: sum(b["counts"][metric] for b in batches) for metric in ("errors", "warnings", "changes")}
    result = {
        "summary": f"{len(batches)}개 파일 중 검사 완료 {len(batches) - failed}개 · 실패/미검증 {failed}개. 오류 {counts['errors']}건 · 검토 {counts['warnings']}건 · 변경 {counts['changes']}건. Excel 재계산 미실행.",
        "status": "failed" if failed == len(batches) else "partial" if failed else "completed",
        "verification_status": "unverified" if any(b["verification_status"] == "unverified" for b in batches) else "findings" if findings else "within_supported_static_scope",
        "batches": batches, "batch_counts": {"total_files": len(batches), "completed_files": len(batches) - failed, "failed_files": failed, "unverified_files": sum(b["verification_status"] == "unverified" for b in batches)},
        "findings": findings, "changes": changes, "counts": counts, "coverage": coverage,
        "supported_schema": all(b["supported_schema"] for b in batches),
        "formula_scope": formula_scope, "formula_scope_supported": all(b["formula_scope_supported"] for b in batches),
        "support_coverage": {"expected_business_headers": list(ALIASES), "accepted_key_headers": sorted(ALIASES["id"]), "files_inspected": len(batches) - failed, "files_failed": failed, "calculation_verified": False},
        "recalculated": False, "rule_version": RULE_VERSION, "baseline_filename": baseline.name,
        "baseline_sha256": next((b["baseline_sha256"] for b in batches if b["baseline_sha256"]), None),
        "repairable_count": 0, "batch_repair_supported": False,
        "suggested_repairs": sum(b["repairable_count"] for b in batches),
        "limitations": [
            "각 비교 파일을 동일한 기준 파일과 독립적으로 검사했습니다. 파일끼리 병합하거나 비교하지 않았습니다.",
            "실패한 파일은 미검증이며 변경·수식 개수를 집계하지 못했습니다. 파일별 상세 결과에서 실패 원인을 확인하세요.",
            "Excel 재계산을 실행하지 않았습니다. 같은 행의 지원 수식과 명시된 업무 스키마에 한정한 정적 검사입니다.",
            "일괄 자동 복원은 지원하지 않습니다. 파일별 결과를 검토한 뒤 해당 파일을 단일 검사로 열어 복원을 진행하세요.",
        ],
    }
    _write_report(result, output_dir, artifacts)
    return result


def repair_workbook(baseline: Path, current: Path, output_dir: Path) -> dict:
    baseline, current, output_dir = map(Path, (baseline, current, output_dir))
    output_dir.mkdir(parents=True, exist_ok=True)
    result = inspect_workbook(baseline, current, output_dir / "before_repair")
    output = output_dir / "repaired.xlsx"
    if output.resolve() in {baseline.resolve(), current.resolve()}:
        raise ValueError("복원 결과가 원본 경로와 같습니다. 다른 결과 폴더를 지정하세요.")
    # openpyxl does not preserve every Excel extension. Reject rich/unsupported packages.
    with ZipFile(current) as archive:
        unsupported = [n for n in archive.namelist() if n.startswith(("xl/slicers/", "xl/pivotTables/", "xl/activeX/", "xl/embeddings/", "xl/threadedComments/", "customXml/", "xl/drawings/", "xl/charts/", "xl/externalLinks/", "xl/queryTables/", "xl/connections.xml", "xl/model/", "xl/vbaProject", "xl/richData/", "_xmlsignatures/"))]
    if unsupported:
        raise ValueError("차트·외부 연결·피벗 등 보존을 보장할 수 없는 기능이 있어 복원하지 않았습니다. 검사 결과만 이용하세요.")
    book = _load(current)
    applied = []
    try:
        for finding in result["findings"]:
            if finding["repairable"]:
                cell = book[finding["sheet"]][finding["cell"]]
                if cell.value != finding["after"]:
                    raise ValueError("검사 후 셀이 달라져 복원을 중단했습니다.")
                cell.value = finding["suggested"]
                applied.append({"sheet": finding["sheet"], "cell": finding["cell"], "before": finding["after"], "after": finding["suggested"], "rule": finding["rule"]})
        if applied:
            book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True, calcMode="auto")
            book.save(output)
        else:
            shutil.copy2(current, output)
    finally:
        book.close()
    remaining = inspect_workbook(baseline, output, output_dir)
    repair = {"filename": output.name, "applied": applied, "applied_count": len(applied), "remaining": remaining, "recalculated": False, "summary": f"기준 수식 {len(applied)}개를 복사본에 복원했습니다. Excel에서 열어 재계산 후 검토하세요.", "original_sha256": result["current_sha256"], "repaired_sha256": _hash(output), "artifacts": [{"name": output.name, "kind": "xlsx"}, {"name": "repair_log.json", "kind": "json"}, *remaining["artifacts"]]}
    (output_dir / "repair_log.json").write_text(json.dumps(repair, ensure_ascii=False, indent=2), encoding="utf-8")
    return repair
